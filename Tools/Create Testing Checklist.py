"""Generate the accessible release checklist and equivalent plain text copy."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'Documentation'
APP = ROOT / 'Application' / 'Accessible Ebook Library Manager.exe'

SECTIONS = [
('Start safely', [
('Launch the correct build', [
 'Close any older copy of the manager and open the Accessible Ebook Library Manager shortcut in the project folder.',
 'Start your usual screen reader and arrow through several books. Open and close a menu with the keyboard.'
], 'Your existing library is present. Book titles and menus are spoken without diagnostic messages, crashes, or an unexpected console window.'),
('Make a backup before changing books', [
 'Open Settings, Settings, then the Library Backup tab. Choose Backup Folder and select a separate folder outside the managed Books folder.',
 'Reopen Settings and choose Back Up Now. After it finishes, use Show Backup Status and check the date and destination.'
], 'The backup completes and includes both the library database and imported book files. Do not continue with repair or duplicate tests if backup fails.'),
]),
('Navigation and feedback', [
('Check the simplified settings', [
 'Open Settings, Settings. Use Ctrl+Tab and Ctrl+Shift+Tab to move through every category, then Tab and Shift+Tab through its controls.',
 'Check General for File Explorer commands, Book List Reading for sound choices, and Library Backup for one Choose Backup Folder command.'
], 'There is no scroll-area keyboard stop. Labels and button roles are spoken. All controls can be reached, and Escape closes Settings.'),
('Read and navigate one list', [
 'Press Ctrl+1. Use Up, Down, Home, and End, then type the first letter of a title.',
 'Check a very long title and titles beginning with A or The. Choose Title and Author in Book List Reading.'
], 'Each move reads only the current book, with its full title. First-letter navigation and title sorting ignore leading A and The. Selected spoken details are respected.'),
('Move between alphabetical columns', [
 'Press Ctrl+2 and use Left and Right to change columns; use Up and Down within each column.',
 'Repeat with Ctrl+3, including an empty letter column and a long title. Return with Ctrl+1.'
], 'The column is identified when entering it, not on every book. Speech does not repeat a title or read the whole column. Arrow navigation stays responsive.'),
('Review recent messages', [
 'Perform several harmless actions such as changing sort order. Open Help, Review Recent Messages.',
 'Use Up and Down to review messages, Home and End to reach the limits, then Tab to Close.'
], 'Up to ten messages from this session are individually readable, newest first. Earlier messages remain available after the status bar clears. Optional announcements in Settings, Feedback are off unless enabled.'),
]),
('Metadata editing', [
('Read individual metadata fields', [
 'On a test book, press Alt+1, Alt+2, Alt+3, Alt+4, and Alt+5 once each.',
 'Press one of those shortcuts twice quickly. Check which field receives focus, then close without saving.'
], 'A single press reads the requested field; it does not unnecessarily repeat the title. A double press opens editing at that field. Other screen-reader commands still work.'),
('Edit and save metadata', [
 'Press F2 on a test book. Use Tab and Shift+Tab through every field, Save, and Cancel. Check field names, existing text, and typing feedback.',
 'Make a small reversible change and save with Ctrl+S. Reopen to verify it. Repeat using the Save button and Alt+S, then restore the original value.'
], 'Each route saves. Focus returns to the book just edited, even if its sort position changes. The correct book is edited when other books are marked.'),
('Recover from a blank title', [
 'Open metadata on a test book, change its author, clear Title, and press Ctrl+S.',
 'Dismiss the warning, verify the author edit remains, and enter a title before saving, or discard the test changes.'
], 'The editor stays open. After the warning, focus is in Title. Other entered values are retained rather than lost.'),
('Confirm discarding changes', [
 'Change a metadata field and press Escape. Listen to the confirmation, then activate Keep Editing.',
 'Press Escape again, Tab to Discard Changes, and activate it. Repeat once with Bulk Edit on test books.'
], 'The confirmation identifies unsaved changes without repeating a question several times. Keep Editing is the safe default. Discard leaves stored metadata unchanged.'),
]),
('Selection search and imports', [
('Select and deselect multiple books', [
 'Mark two test books with Ctrl+Space. Move away and back, then open metadata on a different current book.',
 'Close the editor. Use Ctrl+A to select all visible books, then Ctrl+Shift+A to deselect them.'
], 'Selection is identifiable without moving focus to another book. Individual editing uses the current book. Batch actions use marked books, and Deselect All clears every mark.'),
('Search and return to the library', [
 'Press Ctrl+F, enter part of a known title, and press Enter once. Read the results, then press Escape.',
 'Repeat with a phrase known to be inside a book after Search, Re-index Book Content has completed. Try a query that matches nothing.'
], 'One search is enough. Matching book titles are spoken during navigation. Escape restores the full list. Empty results are understandable and recoverable; report any Not Responding period.'),
('Import books and multipart folders', [
 'Use copies of an EPUB, a searchable PDF, and a DOCX you own. Import them individually, then import a small folder.',
 'On a separate multipart example, choose the one-book combining option and inspect the resulting EPUB.'
], 'Titles and authors are reasonable. ISBN is never used as Source. Front matter comes before numbered chapters; combined parts are in order and originals remain untouched. Record any bad paragraph boundaries.'),
]),
('Metadata discovery and organization', [
('Find missing metadata', [
 'Choose a test book with known missing information and run Book, Find Missing Metadata Wizard.',
 'Review local results, Open Library, and Google Books in order. Cancel once before saving; then repeat and save only information you can verify.'
], 'Progress and results are readable. Cancel applies no proposed edits. Existing values and the book source are preserved. Unavailable information stays blank; incorrect online matches are not silently accepted.'),
('Check the tag view', [
 'Press F2 on a test book and open Tags. Use arrows and Space to select Fiction and Textbooks. Verify checked states are spoken. Press Escape to close the dropdown, then Ctrl+S to save.',
 'Press Ctrl+4. Check both category columns, Other for existing custom tags, and Untagged for a book with none. Reopen the editor and verify both selected tags persisted.'
], 'Ten predefined categories plus Untagged keep navigation bounded. Multiple tags save correctly. Existing custom tags remain stored, and every book stays reachable. Escape closes only the open dropdown, without discarding edits.'),
('Check the missing tags sound', [
 'In Settings, Book List Reading, choose Missing Tags Only. Arrow between a tagged book and one with no tags, then return to the untagged book.',
 'Choose Missing Author or Tags and check a missing-author book too. Turn the sound Off and repeat.'
], 'Missing Tags Only alerts on untagged books, including those with authors. Each new visit sounds once. Missing Author or Tags covers either omission; Off stops all metadata alerts.'),
]),
('Library maintenance', [
('Watch a folder without duplicate imports', [
 'Add a temporary folder under Settings, Watched Folders. Put a copy of a test book there and run Scan Watched Folders Now.',
 'Scan again without changes. During an automatic scan, navigate or edit another book; finally remove the temporary watched folder.'
], 'The first new book appears, unchanged rescans do not duplicate it, and helper notice files are ignored. Report any focus jump or interrupted editing during background completion; this remains an important regression check.'),
('Review duplicate cleanup', [
 'After making a backup, use known duplicate test entries, ideally EPUB and PDF copies of the same edition.',
 'Review Organize, Remove Duplicate Books, Prefer EPUB. Cancel if the proposed matches include different editions or unrelated books.'
], 'Only genuine duplicates are removed from the list, EPUB is preferred as intended, and source files are retained. Do not approve uncertain matches merely to finish this test.'),
('Repair a disposable EPUB copy', [
 'On an imported test copy, use Book, Repair to detect page labels, rebuild the table of contents, and test text cleanup.',
 'Open the repaired EPUB in a reader and compare page navigation, chapter order, and several paragraphs with the original.'
], 'Reports are readable, repair backups exist, and content is not lost. Roman-numbered front matter and numbered pages navigate correctly when supported by the source. Leave backup removal until the repair is verified.'),
]),
('Exports and connected readers', [
('Export original and EPUB copies', [
 'Mark two test books, press Ctrl+E, and export original-format copies into an empty test folder.',
 'Repeat using EPUB output and open the results in a reader. For non-EPUB conversion, Calibre must be installed.'
], 'Both selected books export, existing files are not silently overwritten, and originals remain in the library. Missing conversion tools produce an understandable message.'),
('Export BRF and eBraille', [
 'Export a short known EPUB as BRF, review its report, then repeat as eBraille.',
 'When compatible hardware or software is available, read the opening text and navigate through the output. Check a book with a table or image for warnings.'
], 'The built-in converter produces readable output and accessible warnings. BRF and eBraille remain separate choices. Record device, braille code, and observed reading problems; mark hardware validation Not tested when unavailable.'),
('Send several books to a destination', [
 'Test the destinations you actually use: Voice Dream, Dolphin EasyReader, NLS eReader, or Kindle. Use two marked test books and a test folder first.',
 'If testing Kindle, verify the destination addresses before confirming an actual send. Test HumanWare MTP only when the device is connected.'
], 'The dialog identifies the destination and selected books, and both books arrive. Cancellation sends nothing further. Device-specific and MTP compatibility are not assumed; record the exact model and any failure.'),
]),
('Cancellation recovery and final checks', [
('Cancel work and close safely', [
 'Start a multi-book import, copy, or export on test files. Activate Cancel while it is working.',
 'Start another test operation and close the main window. Choose not to close first; repeat and confirm closing.'
], 'Cancel waits for a safe point, reports cancellation, and keeps completed work. Closing stays responsive while current steps finish. Backup, restore, and individual repair writes may need to finish rather than stop immediately.'),
('Restart and check persistence', [
 'Save a reversible metadata change, choose a display mode and speech setting, then close normally.',
 'Reopen the application from the project shortcut and inspect those settings and the changed book.'
], 'Saved data and supported remembered settings persist. The library opens without a traceback or duplicate import. Restore the original test metadata afterward.'),
('Repeat the accessibility core', [
 'Repeat navigation, metadata editing, selection, row layout, and the new announcement tests with JAWS, NVDA, and Narrator, one screen reader at a time. Mark any unavailable reader Not tested.',
 'If available, check a braille display, high contrast, and your usual Windows text scaling. Ask your independent tester to complete the same core without coaching.'
], 'Names, roles, values, focus, and keyboard access are understandable without mouse use or supplementary speech. Automated tests alone do not prove screen-reader or braille accessibility.'),
('Plan an isolated restore test', [
 'Record the successful backup location and whether it contains both database and book files.',
 'Ask me to prepare an isolated test library before exercising Restore From Backup or damaged-file recovery. Do not restore over your working library just to complete this checklist.'
], 'Recovery is marked Not tested until it is demonstrated safely. Final approval should distinguish completed tests from skipped device, braille, and recovery checks.'),
]),
]

SECTIONS[1:1] = [
('Check the latest speech changes', [
('Hear routine feedback without an OK dialog', [
 'After backing up, change a missing-metadata sound setting and the automatic backup schedule. Return each to your preferred value.',
 'Run Back Up Now and export one test book to an empty folder. Listen when each operation finishes, then open Help, Review Recent Messages.'
], 'Clean completion gives one brief screen-reader notification while the main window is active. No extra OK button is required. Details remain in Recent Messages. Errors and incomplete results still need attention.'),
('Find a cover without extra confirmation', [
 'On a test book, choose the command to look up its cover online. Complete one successful lookup when an internet connection is available.',
 'Check that there is no preliminary online warning and no success OK dialog. Revisit the book and check its cover visually when possible.'
], 'The lookup starts directly and its result is announced briefly. Current-book focus does not jump. A network failure still produces an understandable error; no-match results do not require acknowledgement.'),
('Review online metadata in one dialog', [
 'Write down a test book author and Source, then use Book to look up metadata online. In Review Online Metadata, inspect Fill Blank Fields and the proposed values.',
 'Choose Replace Existing. Correct one suggested author manually, switch modes again, and verify your correction remains. Cancel this trial, then repeat with verified values and save.'
], 'There is no separate fill-or-replace popup. Existing values survive Fill Blank Fields. Replace Existing proposes available online values. Manual corrections survive mode changes. Cancel changes nothing; Save persists accepted values. Book Source is never replaced by a metadata service or ISBN.'),
('Keep background work from interrupting editing', [
 'Use a temporary watched folder with a small test book and allow an automatic scan to run while the metadata editor is open on another book.',
 'Type a reversible edit and navigate the fields. After closing the editor, inspect Recent Messages and the imported book.'
], 'Background results do not take focus or interrupt a modal editor with routine speech. Scan details remain available. An unchanged rescan must not create duplicates.'),
]),
('Check prompts and message history', [
('Keep useful warnings and safe cancellation', [
 'On disposable books, change metadata and press Escape. Choose Keep Editing, then repeat and discard the change.',
 'Try exporting a test book whose stored test file you have temporarily renamed, or cancel a longer operation. Restore the original test filename afterward.'
], 'Unsaved edits still require a decision. Errors, skipped files, partial results, and conversion limitations remain accessible. Clean cancellation uses brief feedback and never claims unfinished work succeeded.'),
('Read messages without losing keyboard focus', [
 'Trigger a routine setting change with supplementary announcements off, then repeat with them on. Listen for duplicate speech and restore your preference.',
 'Generate more than ten harmless status messages and open Help, Review Recent Messages. Read the items with arrows and close the dialog.'
], 'Routine feedback is spoken once in either setting. History retains the latest ten messages, newest first. Closing returns to a usable control. Report missing speech separately for each screen reader.'),
]),
]
SECTIONS.insert(5, ('Rows and bulk editing', [
('Navigate categories as rows', [
 'Choose Organize, Category Layout, Rows. Test Ctrl+2, Ctrl+3, and Ctrl+4. Use Up and Down between categories and Left and Right between books.',
 'Check empty categories and long titles. Switch to Columns, check Ctrl+1, and restart once to verify the remembered layout.'
], 'Rows reverse grouped navigation axes only. Single-list navigation stays Up and Down. Current book and selections survive layout changes. Only the current book is spoken, with category identification on entry.'),
('Apply bulk changes to marked books only', [
 'Record original metadata for two disposable books. Mark them with Ctrl+Space, leave a third unmarked, and open Bulk Edit Metadata.',
 'Check that ISBN, series number, and title are absent. Enable Tags, select two categories, and Apply. Review the confirmation before approving. Reopen all three books and restore test values afterward.'
], 'Only the marked books change. Unique identifiers remain intact. Apply closes after saving without an extra success OK dialog. Tab reaches the final action; Ctrl+S and Alt+S also work.'),
]))

def create():
    OUTPUT.mkdir(exist_ok=True)
    doc=Document()
    sec=doc.sections[0]
    sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.65)
    sec.left_margin=sec.right_margin=Inches(.75)
    for name in ['Normal','Title','Heading 1','Heading 2','List Number']:
        style=doc.styles[name]; style.font.name='Arial'; style.font.color.rgb=RGBColor(0,0,0)
    doc.styles['Normal'].font.size=Pt(11)
    doc.styles['Normal'].paragraph_format.space_after=Pt(5)
    doc.styles['Normal'].paragraph_format.line_spacing=1.05
    doc.styles['Title'].font.size=Pt(20)
    for border in doc.styles['Title'].element.xpath('./w:pPr/w:pBdr'):
        border.getparent().remove(border)
    doc.styles['Heading 1'].font.size=Pt(16)
    doc.styles['Heading 2'].font.size=Pt(12)
    doc.styles['Heading 2'].paragraph_format.space_before=Pt(9)
    doc.styles['Heading 2'].paragraph_format.space_after=Pt(4)
    doc.core_properties.title='Accessible Ebook Library Manager Release Testing Checklist'
    doc.core_properties.subject='Ordered keyboard and accessibility acceptance tests'
    doc.add_paragraph('Accessible Ebook Library Manager\nRelease Testing Checklist', 'Title')
    intro='Work through these tests in order using your screen reader and keyboard. Record Pass, Fail, or Not tested for each numbered test. Stop if books disappear, edits are lost, or files appear damaged. This checklist is for the September 27 2026 release candidate, not a claim that every feature has been verified.'
    doc.add_paragraph(intro)
    doc.add_paragraph('Use test copies of books, not your only originals. The live library remains in AppData. Keep a verified backup before importing, repairing, or removing duplicate entries. Online metadata steps send book identifiers to the lookup services.')
    doc.add_paragraph('Record Windows version, screen reader and version, braille device if any, and the test date. For failures, include the test number, exact keys pressed, expected result, actual speech, and whether it happens again. Do not include passwords or private book text.')
    doc.add_paragraph('Application location')
    p=doc.add_paragraph(str(APP)); p.runs[0].font.size=Pt(9)
    text=['ACCESSIBLE EBOOK LIBRARY MANAGER RELEASE TESTING CHECKLIST','',intro,'','Application: '+str(APP),'','Use disposable copies and back up first. Record screen reader and version, date, Pass/Fail/Not tested, and actual speech. Never test restore over your live library.','']
    case=0
    for section_index,(section,cases) in enumerate(SECTIONS):
        if section_index:
            doc.add_page_break()
        doc.add_paragraph(section,'Heading 1'); text.extend([section.upper(),''])
        for title,steps,expected in cases:
            case+=1
            doc.add_paragraph(f'{case} {title}','Heading 2'); text.append(f'{case}. {title}')
            for step_index,step in enumerate(steps,1):
                p=doc.add_paragraph()
                p.add_run(f'Step {step_index}: ').bold=True
                p.add_run(step)
                text.append(f'   {step_index}. {step}')
            p=doc.add_paragraph(); p.add_run('Expected: ').bold=True; p.add_run(expected)
            p=doc.add_paragraph('Result: Pass / Fail / Not tested. Notes:'); p.paragraph_format.space_after=Pt(7)
            text.extend(['Expected: '+expected,'Result: Pass / Fail / Not tested. Notes:',''])
    doc.save(OUTPUT/'Release Testing Checklist.docx')
    (OUTPUT/'Release Testing Checklist.txt').write_text('\n'.join(text),encoding='utf-8')
    print(f'Created {case} ordered tests.')

if __name__=='__main__':
    create()
